from datetime import datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import httpx
import pytest
import pytest_asyncio
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.orm import selectinload

from app.database import Base, get_db
from app.models import Project, Task, TaskStatus, TaskTimeSession
from app.routers import tasks


@pytest_asyncio.fixture
async def time_api(monkeypatch):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as db:
        db.add(Project(id=1, name="<project-name>"))
        db.add_all([
            TaskStatus(id=4, project_id=1, name="К выполнению", sort_order=0),
            TaskStatus(id=1, project_id=1, name="В работе", sort_order=0),
            TaskStatus(id=2, project_id=1, name="Тестирование", sort_order=1),
            TaskStatus(id=3, project_id=1, name="Деплой", sort_order=2),
        ])
        db.add_all([
            Task(id=1, project_id=1, status_id=2, title="<task-title>", estimated_minutes=30),
            Task(id=2, project_id=1, status_id=2, title="<other-task-title>", estimated_minutes=15),
            Task(id=3, project_id=1, status_id=2, title="<task-without-plan>", manual_actual_seconds=120),
            Task(id=4, project_id=1, status_id=4, title="<unstarted-task>", estimated_minutes=20),
        ])
        await db.commit()

    async def db_dependency():
        async with sessions() as db:
            yield db

    monkeypatch.setattr(tasks, "get_settings", lambda: SimpleNamespace(mcp_api_key="<mcp-api-key>"))
    events = []
    monkeypatch.setattr(tasks, "broadcast", events.append)
    app = FastAPI()
    app.include_router(tasks.global_router)
    app.include_router(tasks.router)
    app.dependency_overrides[get_db] = db_dependency
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="https://<service-host>",
        headers={"Authorization": "Bearer <mcp-api-key>"},
    ) as client:
        yield client, sessions, events
    await engine.dispose()


@pytest.mark.asyncio
@pytest.mark.parametrize(("phase", "source", "target"), [
    ("work", 1, 2),
    ("testing", 2, 3),
    ("work", 1, 3),
    ("testing", 2, 1),
    ("testing", 2, 2),
    ("testing", 1, 2),
])
async def test_manual_time_completion_is_independent_from_kanban_status(time_api, phase, source, target):
    client, sessions, events = time_api
    await client.patch("/api/projects/1/tasks/1", json={"status_id": source})
    await client.post("/api/tasks/1/time/manual/start", json={"phase": phase})
    async with sessions() as db:
        interval = (await db.execute(select(TaskTimeSession))).scalar_one()
        interval.started_at = datetime.utcnow() - timedelta(minutes=2)
        await db.commit()
    paused = await client.post("/api/tasks/1/time/manual/pause")
    assert paused.json()["time_tracking_status"] == "paused"
    events.clear()
    changed = await client.patch("/api/tasks/1/time/status", json={"status_id": target})
    assert changed.status_code == 200
    assert changed.json()["time_tracking_status"] == "paused"
    assert changed.json()["actual_seconds"] == paused.json()["actual_seconds"]
    assert changed.json()["manual_time_phase"] is None
    listed = (await client.get("/api/tasks")).json()
    assert next(task for task in listed if task["id"] == 1)["time_tracking_status"] == "paused"
    if source != target:
        assert events[-1]["type"] == "task_changed"
    resumed = await client.post("/api/tasks/1/time/manual/start", json={"phase": "testing"})
    assert resumed.json()["time_tracking_status"] == "running"
    assert resumed.json()["actual_seconds"] >= paused.json()["actual_seconds"]


@pytest.mark.asyncio
async def test_agent_time_phase_is_independent_from_kanban_status(time_api):
    client, _, _ = time_api
    await client.patch("/api/projects/1/tasks/1", json={"status_id": 1})

    testing = await client.post(
        "/api/tasks/1/time/start",
        json={"worker_id": "<testing-worker>", "phase": "testing"},
    )
    work = await client.post(
        "/api/tasks/2/time/start",
        json={"worker_id": "<work-worker>", "phase": "work"},
    )

    assert testing.status_code == 200
    assert testing.json()["phase"] == "testing"
    assert work.status_code == 200
    assert work.json()["phase"] == "work"
    assert (await client.get("/api/tasks/1")).json()["status_id"] == 1
    assert (await client.get("/api/tasks/2")).json()["status_id"] == 2


@pytest.mark.asyncio
async def test_work_to_testing_status_switch_starts_testing_time_phase(time_api):
    client, sessions, _ = time_api
    changed = await client.patch("/api/projects/1/tasks/1", json={"status_id": 1})
    assert changed.status_code == 200

    started = await client.post(
        "/api/tasks/1/time/start",
        json={"worker_id": "<worker-id>", "phase": "work"},
    )
    assert started.status_code == 200

    changed = await client.patch("/api/projects/1/tasks/1", json={"status_id": 2})
    assert changed.status_code == 200
    assert changed.json()["time_tracking_status"] == "running"
    assert changed.json()["manual_time_phase"] is None

    async with sessions() as db:
        intervals = (
            await db.execute(
                select(TaskTimeSession)
                .where(TaskTimeSession.task_id == 1)
                .order_by(TaskTimeSession.started_at.asc(), TaskTimeSession.id.asc())
            )
        ).scalars().all()
        assert [(interval.phase, interval.ended_at is None) for interval in intervals] == [
            ("work", False),
            ("testing", True),
        ]


@pytest.mark.asyncio
@pytest.mark.parametrize("phase", ["work", "testing"])
async def test_manual_timer_accumulates_selected_phase_and_blocks_status_until_paused(time_api, phase):
    client, sessions, events = time_api
    client.headers.pop("Authorization")
    started = await client.post("/api/tasks/1/time/manual/start", json={"phase": phase})
    assert started.status_code == 200
    assert started.json()["manual_time_phase"] == phase
    assert started.json()["status_id"] == 2
    async with sessions() as db:
        interval = (await db.execute(select(TaskTimeSession))).scalar_one()
        interval.started_at = datetime.utcnow() - timedelta(minutes=3)
        await db.commit()
    assert (await client.post("/api/tasks/1/time/manual/start", json={"phase": phase})).status_code == 409
    assert (await client.patch("/api/tasks/1/time/status", json={"status_id": 1})).status_code == 409
    paused = await client.post("/api/tasks/1/time/manual/pause")
    assert paused.status_code == 200
    assert paused.json()["manual_time_phase"] is None
    assert paused.json()["time_tracking_status"] == "paused"
    assert paused.json()[f"{phase}_seconds"] >= 180
    assert paused.json()["actual_seconds"] == paused.json()[f"{phase}_seconds"]
    assert (await client.patch("/api/tasks/1/time/status", json={"status_id": 999})).status_code == 404
    changed = await client.patch("/api/tasks/1/time/status", json={"status_id": 1})
    assert changed.status_code == 200
    assert changed.json()["time_tracking_status"] == "paused"
    assert changed.json()["status_id"] == 1
    assert events[-1]["type"] == "task_changed"
    resumed = await client.post("/api/tasks/1/time/manual/start", json={"phase": phase})
    assert resumed.json()[f"{phase}_seconds"] >= 180
    assert (await client.get("/api/tasks")).json()[0]["manual_time_phase"] == phase


@pytest.mark.asyncio
async def test_agent_takes_over_manual_timer_and_cannot_be_stopped_from_manual_controls(time_api):
    client, sessions, _ = time_api
    await client.post("/api/tasks/1/time/manual/start", json={"phase": "work"})
    async with sessions() as db:
        interval = (await db.execute(select(TaskTimeSession))).scalar_one()
        interval.started_at = datetime.utcnow() - timedelta(minutes=2)
        await db.commit()
    response = await client.post("/api/tasks/1/time/start", json={"worker_id": "<worker-id>", "phase": "testing"})
    assert response.status_code == 200
    task = (await client.get("/api/tasks/1")).json()
    assert task["manual_time_phase"] is None
    assert task["work_seconds"] >= 120
    assert task["time_tracking_status"] == "running"
    assert (await client.post("/api/tasks/1/time/manual/pause")).status_code == 409
    assert (await client.post("/api/tasks/1/time/manual/start", json={"phase": "work"})).status_code == 409
    assert (await client.patch("/api/tasks/1/time/status", json={"status_id": 1})).status_code == 409
    assert (await client.post("/api/tasks/1/time/pause", json={"worker_id": "manual:task:1"})).status_code == 422
    await client.post("/api/tasks/1/time/pause", json={"worker_id": "<worker-id>", "completed": True})
    assert (await client.post("/api/tasks/1/time/manual/start", json={"phase": "testing"})).status_code == 200


@pytest.mark.asyncio
@pytest.mark.parametrize(("phase", "status_id", "status_name"), [("work", 1, "В работе"), ("testing", 2, "Тестирование")])
async def test_pause_completion_and_resume_are_independent_of_kanban(time_api, phase, status_id, status_name):
    client, sessions, _ = time_api
    changed = await client.patch("/api/projects/1/tasks/1", json={"status_id": status_id})
    assert changed.status_code == 200
    initial = (await client.get("/api/tasks/1")).json()
    assert initial["time_tracking_status"] is None

    body = {"worker_id": "<worker-id>", "phase": phase}
    started = await client.post("/api/tasks/1/time/start", json=body)
    assert started.status_code == 200
    assert started.json()["time_tracking_status"] == "running"
    async with sessions() as db:
        interval = (await db.execute(select(TaskTimeSession))).scalar_one()
        interval.started_at = datetime.utcnow() - timedelta(minutes=3)
        await db.commit()

    paused = await client.post("/api/tasks/1/time/pause", json={"worker_id": "<worker-id>"})
    assert paused.status_code == 200
    assert paused.json()["time_tracking_status"] == "paused"
    before_finish = (await client.get("/api/tasks/1")).json()
    completed_body = {"worker_id": "<worker-id>", "completed": True}
    finished = await client.post("/api/tasks/1/time/pause", json=completed_body)
    assert finished.json()["time_tracking_status"] == "completed"
    after_finish = (await client.get("/api/tasks/1")).json()
    assert after_finish["status"]["name"] == status_name
    assert after_finish[f"{phase}_seconds"] == before_finish[f"{phase}_seconds"]
    assert after_finish["actual_seconds"] >= 180
    assert finished.json()["ended_at"] == paused.json()["ended_at"]

    repeated = await client.post("/api/tasks/1/time/pause", json={"worker_id": "<worker-id>"})
    assert repeated.json()["time_tracking_status"] == "completed"
    resumed = await client.post("/api/tasks/1/time/start", json=body)
    assert resumed.json()["time_tracking_status"] == "running"
    async with sessions() as db:
        task = (await db.execute(select(Task).options(selectinload(Task.time_sessions)).where(Task.id == 1))).scalar_one()
        assert len(task.time_sessions) == 2
        assert task.time_sessions[0].completed is True
        assert task.time_sessions[1].completed is False
        assert task.actual_seconds >= 180


@pytest.mark.asyncio
async def test_parallel_workers_and_switching_tasks_keep_running_status_and_notify_both_tasks(time_api):
    client, _, events = time_api
    for worker in ("<first-worker>", "<second-worker>"):
        response = await client.post("/api/tasks/1/time/start", json={"worker_id": worker, "phase": "testing"})
        assert response.status_code == 200
    finished = await client.post("/api/tasks/1/time/pause", json={"worker_id": "<second-worker>", "completed": True})
    assert finished.json()["active"] is False
    assert finished.json()["time_tracking_status"] == "running"
    events.clear()
    switched = await client.post("/api/tasks/2/time/start", json={"worker_id": "<first-worker>", "phase": "testing"})
    assert switched.status_code == 200
    assert {event["task_id"] for event in events} == {1, 2}
    first = (await client.get("/api/tasks/1")).json()
    assert first["time_tracking_status"] == "paused"
    assert (await client.get("/api/tasks/2")).json()["time_tracking_status"] == "running"


@pytest.mark.asyncio
@pytest.mark.parametrize("project_id", [None, 1])
async def test_time_status_is_in_list_but_not_exported_to_xlsx(time_api, project_id):
    client, _, _ = time_api
    await client.post("/api/tasks/1/time/start", json={"worker_id": "<worker-id>", "phase": "testing"})
    await client.post("/api/tasks/1/time/pause", json={"worker_id": "<worker-id>", "completed": True})
    url = "/api/tasks" if project_id is None else f"/api/projects/{project_id}/tasks"
    listed = (await client.get(url)).json()
    assert next(task for task in listed if task["id"] == 1)["time_tracking_status"] == "completed"
    exported = await client.post("/api/tasks/time/export/xlsx", json={"task_ids": [1, 2], "project_id": project_id})
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(2, headers.index("Статус") + 1).value == "Тестирование"
    assert "Статус времени" not in headers
    assert sheet.cell(4, headers.index("План") + 1).value == "45 мин"
    assert sheet.cell(4, headers.index("Факт") + 1).value == "0 мин"
    assert sheet.max_column == (9 if project_id is None else 8)


@pytest.mark.asyncio
async def test_time_deviation_uses_zero_for_missing_plan_and_skips_unstarted_plan(time_api):
    client, _, _ = time_api
    exported = await client.post(
        "/api/tasks/time/export/xlsx",
        json={"task_ids": [1, 2, 3, 4], "project_id": 1},
    )
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content)).active
    headers = [cell.value for cell in sheet[1]]
    deviation_column = headers.index("Отклонение") + 1
    rows = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row)}

    assert sheet.cell(rows[1], deviation_column).value == "−30 мин"
    assert sheet.cell(rows[3], deviation_column).value == "+2 мин"
    assert sheet.cell(rows[4], deviation_column).value == "—"
    assert sheet.cell(sheet.max_row, deviation_column).value == "−43 мин"
